import os
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

def render_pptx_slides(pptx_path, output_dir):
    os.makedirs(output_dir, exist_ok=True)
    prs = Presentation(pptx_path)
    
    slide_w_in = prs.slide_width.inches
    slide_h_in = prs.slide_height.inches
    
    scale = 1920.0 / slide_w_in
    target_w = int(round(slide_w_in * scale))
    target_h = int(round(slide_h_in * scale))
    
    print(f"Rendering {len(prs.slides)} slides at {target_w}x{target_h} (scale={scale:.2f} px/in)...")
    
    font_paths = [
        "C:\\Windows\\Fonts\\segoeui.ttf",
        "C:\\Windows\\Fonts\\arial.ttf",
        "C:\\Windows\\Fonts\\calibri.ttf"
    ]
    font_bold_paths = [
        "C:\\Windows\\Fonts\\segoeuib.ttf",
        "C:\\Windows\\Fonts\\arialbd.ttf",
        "C:\\Windows\\Fonts\\calibrib.ttf"
    ]
    
    regular_font_path = next((p for p in font_paths if os.path.exists(p)), None)
    bold_font_path = next((p for p in font_bold_paths if os.path.exists(p)), regular_font_path)
    
    def get_font(size_pt, bold=False):
        size_px = int(round(size_pt * 4 / 3 * (scale / 96.0)))
        size_px = max(size_px, 10)
        try:
            f_path = bold_font_path if bold else regular_font_path
            if f_path:
                return ImageFont.truetype(f_path, size_px)
        except Exception:
            pass
        return ImageFont.load_default()

    def draw_wrapped_text(draw, text, x, y, max_w, font, fill):
        # Split on explicit newlines first
        lines = text.split("\n")
        cur_y = y
        line_h = font.getbbox("Ay")[3] - font.getbbox("Ay")[1] + 5

        for line in lines:
            words = line.split(" ")
            current_line = ""
            for word in words:
                test_line = f"{current_line} {word}".strip()
                bbox = font.getbbox(test_line)
                w = bbox[2] - bbox[0]
                if w <= max_w:
                    current_line = test_line
                else:
                    if current_line:
                        draw.text((x, cur_y), current_line, font=font, fill=fill)
                        cur_y += line_h
                    current_line = word
            if current_line:
                draw.text((x, cur_y), current_line, font=font, fill=fill)
                cur_y += line_h
        return cur_y

    for slide_idx, slide in enumerate(prs.slides, 1):
        bg_rgb = (250, 250, 247) # default warm ivory
        if slide_idx == 1:
            bg_rgb = (11, 25, 44) # Slide 1 Navy
            
        img = Image.new("RGB", (target_w, target_h), bg_rgb)
        draw = ImageDraw.Draw(img)
        
        for shape in slide.shapes:
            left = int(round(shape.left.inches * scale))
            top = int(round(shape.top.inches * scale))
            width = int(round(shape.width.inches * scale))
            height = int(round(shape.height.inches * scale))
            right = left + width
            bottom = top + height
            
            # 1. Picture shape
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image_bytes = shape.image.blob
                    from io import BytesIO
                    pic_img = Image.open(BytesIO(image_bytes))
                    pic_resized = pic_img.resize((width, height), Image.Resampling.LANCZOS)
                    img.paste(pic_resized, (left, top))
                except Exception as e:
                    print(f"Error drawing picture on slide {slide_idx}: {e}")
                    
            # 2. Auto Shape
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill_color = None
                line_color = None
                
                try:
                    if shape.fill.type and shape.fill.fore_color.rgb:
                        c = shape.fill.fore_color.rgb
                        fill_color = (c[0], c[1], c[2])
                except Exception:
                    pass
                    
                try:
                    if shape.line and shape.line.color and shape.line.color.rgb:
                        c = shape.line.color.rgb
                        line_color = (c[0], c[1], c[2])
                except Exception:
                    pass
                    
                if fill_color:
                    draw.rounded_rectangle([left, top, right, bottom], radius=8, fill=fill_color, outline=line_color, width=2 if line_color else 0)
                elif line_color:
                    draw.rounded_rectangle([left, top, right, bottom], radius=8, outline=line_color, width=2)
                    
                # Text inside shape
                if shape.has_text_frame:
                    tf = shape.text_frame
                    margin_l = int(round(tf.margin_left.inches * scale)) if tf.margin_left else 15
                    margin_t = int(round(tf.margin_top.inches * scale)) if tf.margin_top else 12
                    margin_r = int(round(tf.margin_right.inches * scale)) if tf.margin_right else 15
                    cur_y = top + margin_t
                    max_text_w = width - margin_l - margin_r
                    
                    for p in tf.paragraphs:
                        if not p.text.strip():
                            cur_y += 8
                            continue
                        font_size_pt = p.font.size.pt if p.font and p.font.size else 10
                        is_bold = bool(p.font and p.font.bold)
                        text_color = (26, 26, 26)
                        try:
                            if p.font and p.font.color and p.font.color.rgb:
                                c = p.font.color.rgb
                                text_color = (c[0], c[1], c[2])
                        except Exception:
                            pass
                            
                        font = get_font(font_size_pt, is_bold)
                        
                        # Check center alignment
                        from pptx.enum.text import PP_ALIGN
                        is_center = (p.alignment == PP_ALIGN.CENTER)
                        if is_center:
                            bbox = font.getbbox(p.text)
                            tw = bbox[2] - bbox[0]
                            cx = left + (width - tw) // 2
                            draw.text((cx, cur_y), p.text, font=font, fill=text_color)
                            line_h = font.getbbox("Ay")[3] - font.getbbox("Ay")[1] + 5
                            cur_y += line_h
                        else:
                            cur_y = draw_wrapped_text(draw, p.text, left + margin_l, cur_y, max_text_w, font, text_color)
                        
            # 3. Textbox
            elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                tf = shape.text_frame
                cur_y = top
                max_text_w = width
                for p in tf.paragraphs:
                    if not p.text.strip():
                        cur_y += 8
                        continue
                    font_size_pt = p.font.size.pt if p.font and p.font.size else 11
                    is_bold = bool(p.font and p.font.bold)
                    text_color = (26, 26, 26)
                    try:
                        if p.font and p.font.color and p.font.color.rgb:
                            c = p.font.color.rgb
                            text_color = (c[0], c[1], c[2])
                    except Exception:
                        pass
                        
                    font = get_font(font_size_pt, is_bold)
                    from pptx.enum.text import PP_ALIGN
                    if p.alignment == PP_ALIGN.CENTER:
                        bbox = font.getbbox(p.text)
                        tw = bbox[2] - bbox[0]
                        cx = left + (width - tw) // 2
                        draw.text((cx, cur_y), p.text, font=font, fill=text_color)
                        line_h = font.getbbox("Ay")[3] - font.getbbox("Ay")[1] + 5
                        cur_y += line_h
                    elif p.alignment == PP_ALIGN.RIGHT:
                        bbox = font.getbbox(p.text)
                        tw = bbox[2] - bbox[0]
                        rx = right - tw
                        draw.text((rx, cur_y), p.text, font=font, fill=text_color)
                        line_h = font.getbbox("Ay")[3] - font.getbbox("Ay")[1] + 5
                        cur_y += line_h
                    else:
                        cur_y = draw_wrapped_text(draw, p.text, left, cur_y, max_text_w, font, text_color)
                    
        out_fn = os.path.join(output_dir, f"slide_{slide_idx:02d}.png")
        img.save(out_fn, "PNG")
        print(f"Saved preview: {out_fn}")

if __name__ == "__main__":
    pptx_file = sys.argv[1] if len(sys.argv) > 1 else os.path.join("docs", "presentation", "Niventra_Capital_Advisory_Website_Overview.pptx")
    out_folder = os.path.join("docs", "presentation", "previews", "overview")
    render_pptx_slides(pptx_file, out_folder)
